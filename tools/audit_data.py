"""Audit supplied measurements without modifying the workbook or training a model."""
from pathlib import Path
import csv
import math
import openpyxl

ROOT = Path(__file__).resolve().parents[1]

def numeric_rows(sheet):
    result = []
    invalid = []
    for index, row in enumerate(sheet.values, 1):
        if index == 1 or not any(v is not None for v in row):
            continue
        if len(row) == 8 and all(isinstance(v, (int, float)) and math.isfinite(v) for v in row):
            result.append((index, tuple(row)))
        else:
            invalid.append(index)
    return result, invalid

def save_csv(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', newline='', encoding='utf-8') as stream:
        writer = csv.DictWriter(stream, fieldnames=rows[0])
        writer.writeheader()
        writer.writerows(rows)

def main():
    workbook = openpyxl.load_workbook(ROOT / 'data/raw/HSR-P.xlsx', data_only=True)
    raw = {}
    audit = []
    for sheet in workbook:
        rows, invalid = numeric_rows(sheet)
        raw[sheet.title] = rows
        audit.append(dict(sheet=sheet.title, worksheet_max_row=sheet.max_row,
                          numeric_records=len(rows), unique_records=len(set(v for _, v in rows)),
                          invalid_nonempty_rows=len(invalid)))
    lookup = {}
    for label, rows in raw.items():
        for index, values in rows:
            lookup.setdefault(values, []).append((label, index))
    supplied = openpyxl.load_workbook(ROOT / 'data/provided_selection/HSR-P_DE_top100.xlsx', data_only=True)
    matches = []
    for sheet in supplied:
        rows, invalid = numeric_rows(sheet)
        if invalid:
            raise ValueError(f'Invalid selection rows: {sheet.title}: {invalid}')
        for index, values in rows:
            sources = lookup.get(values, [])
            matches.append(dict(provided_sheet=sheet.title, provided_excel_row=index,
                                matching_raw_sheets=';'.join(k for k, _ in sources),
                                matching_raw_rows=';'.join(str(i) for _, i in sources),
                                same_sheet_match=any(k == sheet.title for k, _ in sources)))
    destination = ROOT / 'results/audit'
    save_csv(destination / 'dataset_audit.csv', audit)
    save_csv(destination / 'provided_selection_matches.csv', matches)
    for entry in audit:
        print(f"{entry['sheet']}: {entry['numeric_records']} numeric records")
    mismatches = sum(not row['same_sheet_match'] for row in matches)
    print(f'Provided selection: {mismatches}/{len(matches)} records do not match their labelled raw sheet.')
    print('No measurements were altered. MATLAB calibration was not executed.')

if __name__ == '__main__':
    main()
